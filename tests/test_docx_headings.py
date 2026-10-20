import io
from copy import deepcopy
import pytest
from docx import Document
from lxml import etree as ET
from workshop.office.engine import OfficeEngine
from workshop.office.package import Package, PackageError, q
from workshop.office.repairs import plan_repairs, apply_repairs
from workshop.contracts import digest
from tests.test_office import save


def apply(data, roles):
    objects = [
        o
        for o in OfficeEngine().inspect(data, "Headings.docx")["objects"]
        if o["kind"] == "paragraph" and o["target"]["part"] == "word/document.xml"
    ]
    plan = plan_repairs(
        data,
        "Headings.docx",
        digest(data),
        [
            dict(object_id=objects[i]["id"], action="set_paragraph_role", value=role)
            for i, role in roles
        ],
        "heading-test-key",
    )
    return apply_repairs(data, plan)


def contents(node):
    node = deepcopy(node)
    props = node.find(q("w", "pPr"))
    if props is not None:
        node.remove(props)
    return ET.tostring(node, method="c14n")


@pytest.mark.parametrize(
    "role", ["body", "title"] + [f"heading{i}" for i in range(1, 10)]
)
def test_roles_preserve_runs_and_unrelated_parts(role):
    doc = Document()
    p = doc.add_heading("Existing ", 2)
    p.add_run("{{ users[0].name }}").bold = True
    p.add_run(" unchanged").italic = True
    p.paragraph_format.keep_together = True
    doc.add_paragraph("Other paragraph")
    data = save(doc)
    after, report, receipt = apply(data, [(0, role)])
    before_pkg, after_pkg = Package(data), Package(after)
    before_p = next(before_pkg.roots[before_pkg.main].iter(q("w", "p")))
    after_p = next(after_pkg.roots[after_pkg.main].iter(q("w", "p")))
    assert contents(before_p) == contents(after_p)
    assert Document(io.BytesIO(after)).paragraphs[0].paragraph_format.keep_together
    assert (
        next(o for o in report["objects"] if o["kind"] == "paragraph")["role"] == role
    )
    assert receipt["verification"]["operation_idempotence"] == "pass"
    for name in before_pkg.members:
        if name not in receipt["changed_parts"]:
            assert before_pkg.members[name] == after_pkg.members[name]


def test_missing_style_and_duplicate_text_targets():
    doc = Document()
    doc.styles["Heading 9"].delete()
    doc.add_paragraph("Same text")
    doc.add_paragraph("Same text")
    after, report, _ = apply(save(doc), [(1, "heading9")])
    paras = [o for o in report["objects"] if o["kind"] == "paragraph"]
    assert [o["role"] for o in paras] == ["body", "heading9"]
    assert Document(io.BytesIO(after)).paragraphs[1].style.name == "Heading 9"


def test_bulk_roles_and_reject_arbitrary_text():
    doc = Document()
    doc.add_paragraph("One")
    doc.add_paragraph("Two")
    data = save(doc)
    after, report, _ = apply(data, [(0, "heading1"), (1, "heading2")])
    assert [p.text for p in Document(io.BytesIO(after)).paragraphs] == ["One", "Two"]
    assert [o["role"] for o in report["objects"] if o["kind"] == "paragraph"] == [
        "heading1",
        "heading2",
    ]
    with pytest.raises(PackageError, match="Choose Body"):
        apply(data, [(0, "Replacement text")])


def test_table_name_changes_only_selected_table_properties():
    from tests.test_office import docx_bytes, repair

    data = docx_bytes()
    (after, report, receipt), _ = repair(
        data, "table", "set_table_title", "Applicant contact details"
    )
    before_pkg, after_pkg = Package(data), Package(after)
    tables = [o for o in report["objects"] if o["kind"] == "table"]
    assert [o["table_title"] for o in tables] == ["Applicant contact details", ""]
    before_tables = list(before_pkg.roots[before_pkg.main].iter(q("w", "tbl")))
    after_tables = list(after_pkg.roots[after_pkg.main].iter(q("w", "tbl")))
    assert ET.tostring(before_tables[1], method="c14n") == ET.tostring(
        after_tables[1], method="c14n"
    )
    for before, after_row in zip(
        before_tables[0].findall(q("w", "tr")), after_tables[0].findall(q("w", "tr"))
    ):
        assert ET.tostring(before, method="c14n") == ET.tostring(
            after_row, method="c14n"
        )
    assert receipt["changed_parts"] == ["word/document.xml"]


def test_reused_image_description_only_changes_selected_instance(tmp_path):
    from PIL import Image
    from tests.test_office import repair

    path = tmp_path / "image.png"
    Image.new("RGB", (20, 20), "blue").save(path)
    doc = Document()
    doc.add_picture(str(path))
    doc.add_picture(str(path))
    data = save(doc)
    (after, report, receipt), _ = repair(
        data, "image", "set_alt_text", "Blue indicates completed tasks"
    )
    images = [o for o in report["objects"] if o["kind"] == "image"]
    assert [o["value"] for o in images] == ["Blue indicates completed tasks", ""]
    (decorated, report, _), _ = repair(after, "image", "mark_decorative", "")
    images = [o for o in report["objects"] if o["kind"] == "image"]
    assert images[0]["decorative"] and not images[1]["decorative"]
    assert images[0]["value"] == ""
    for part, payload in Package(data).members.items():
        if part != "word/document.xml":
            assert Package(decorated).members[part] == payload

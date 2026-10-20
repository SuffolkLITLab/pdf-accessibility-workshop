import io
from pathlib import Path
from unittest.mock import patch
import pytest
import pikepdf
from docx import Document
from lxml import etree as ET
from workshop.contracts import digest
from workshop.office.engine import OfficeEngine
from workshop.office.package import Package, PackageError, Conflict, q
from workshop.office.preview import instrument, convert
from workshop.office.repairs import plan_repairs, apply_repairs
from tests.test_office import save


def merge(data, indices):
    report = OfficeEngine().inspect(data, "Blocks.docx")
    paras = [
        o
        for o in report["objects"]
        if o["kind"] == "paragraph" and o["target"]["part"] == "word/document.xml"
    ]
    selected = [paras[i] for i in indices]
    plan = plan_repairs(
        data,
        "Blocks.docx",
        digest(data),
        [
            {
                "object_id": selected[0]["id"],
                "action": "merge_paragraphs",
                "value": [o["id"] for o in selected],
            }
        ],
        "merge-retry-key",
    )
    return apply_repairs(data, plan), plan


def test_merge_retains_runs_properties_and_retry_determinism():
    doc = Document()
    first = doc.add_heading("First {{ value }}", 2)
    second = doc.add_paragraph("Second ")
    second.add_run("bold").bold = True
    doc.add_paragraph("Untouched")
    data = save(doc)
    (after, report, receipt), plan = merge(data, [0, 1])
    reopened = Document(io.BytesIO(after))
    assert len(reopened.paragraphs) == 2
    assert reopened.paragraphs[0].text == "First {{ value }}\nSecond bold"
    assert reopened.paragraphs[0].style.name == "Heading 2"
    assert reopened.paragraphs[0].runs[-1].bold
    assert reopened.paragraphs[1].text == "Untouched"
    assert receipt["verification"]["merge_content_preservation"] == "pass"
    assert apply_repairs(data, plan)[0] == after
    with pytest.raises(Conflict):
        apply_repairs(after, plan)
    for part, payload in Package(data).members.items():
        if part != "word/document.xml":
            assert Package(after).members[part] == payload


def test_merge_rejects_nonadjacent_cells_and_section_boundaries():
    doc = Document()
    doc.add_paragraph("One")
    doc.add_paragraph("Two")
    doc.add_paragraph("Three")
    with pytest.raises(PackageError, match="adjacent"):
        merge(save(doc), [0, 2])
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "One"
    table.cell(0, 1).text = "Two"
    with pytest.raises(PackageError, match="parent"):
        merge(save(doc), [0, 1])
    doc = Document()
    doc.add_paragraph("One")
    doc.add_section()
    doc.add_paragraph("Two")
    with pytest.raises(PackageError, match="ordinary"):
        merge(save(doc), [0, 1])


def test_preview_anchors_are_native_and_only_in_derivative():
    data = Path("tests/semantic-example.docx").read_bytes()
    report = OfficeEngine().inspect(data, "Preview.docx")
    derivative, anchors = instrument(data, report["objects"])
    assert data != derivative
    assert report["blocks"][0]["kind"] == "paragraph"
    assert any(b["kind"] == "table" for b in report["blocks"])
    assert len({a["start"] for a in anchors}) == len(anchors)
    assert {a["kind"] for a in anchors} == {"paragraph", "table", "image"}
    source, modified = Package(data), Package(derivative)
    assert not list(source.roots[source.main].iter(q("w", "bookmarkStart")))
    names = {
        n.get(q("w", "name"))
        for n in modified.roots[modified.main].iter(q("w", "bookmarkStart"))
    }
    assert all(a["start"] in names and a["end"] in names for a in anchors)
    for part in source.members:
        if part != source.main:
            assert source.members[part] == modified.members[part]


def test_conversion_is_bounded_digest_bound_and_not_validation(monkeypatch):
    monkeypatch.setenv("GOTENBERG_URL", "https://conversion.example.invalid")
    data = Path("tests/semantic-example.docx").read_bytes()
    out = io.BytesIO()
    with pikepdf.new() as pdf:
        pdf.add_blank_page()
        pdf.save(out)

    class Response:
        status_code = 200

        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

        def iter_content(self, size):
            yield out.getvalue()

    with patch(
        "workshop.office.preview.requests.post", return_value=Response()
    ) as post:
        with pytest.raises(Conflict):
            convert(data, "Private filename.docx", "stale")
        post.assert_not_called()
        result = convert(data, "Private filename.docx", digest(data))
        assert result["source_sha256"] == digest(data)
        assert result["accessibility_validation"] == "not_run"
        assert post.call_args.kwargs["allow_redirects"] is False
        assert post.call_args.kwargs["files"]["files"][0] == "document.docx"
        assert "pdfua" not in post.call_args.kwargs["data"]
        Response.status_code = 302
        with pytest.raises(PackageError, match="could not convert"):
            convert(data, "Preview.docx", digest(data))

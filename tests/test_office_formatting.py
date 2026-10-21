import io
import json
from pathlib import Path
from unittest.mock import patch
import pytest
from docx import Document
from docx.shared import Pt, RGBColor
from pptx import Presentation
from pptx.util import Inches, Pt as SlidePt
from pptx.dml.color import RGBColor as SlideColor
from lxml import etree as ET
from workshop.contracts import digest
from workshop.office.engine import OfficeEngine
from workshop.office.package import Package, PackageError, q, NS
from workshop.office.repairs import plan_repairs, apply_repairs
from workshop.office.ai_drafts import draft_repairs
from workshop.office.formatting import ratio


def fixture(fmt, size=12, color="999999", bold=False):
    out = io.BytesIO()
    if fmt == "docx":
        doc = Document()
        style = doc.styles["Normal"]
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = bold
        p = doc.add_paragraph()
        p.add_run("Hello {{ users[0].name }} ").italic = True
        p.add_run("readers")
        doc.save(out)
    else:
        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = SlideColor(255, 255, 255)
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        shape.fill.background()
        p = shape.text_frame.paragraphs[0]
        for words in ("Hello {{ users[0].name }} ", "readers"):
            run = p.add_run()
            run.text = words
            run.font.size = SlidePt(size)
            run.font.bold = bold
            run.font.color.rgb = SlideColor.from_string(color)
        p.runs[0].font.italic = True
        doc.save(out)
    return out.getvalue()


def target(data, fmt):
    report = OfficeEngine().inspect(data, "Sample." + fmt)
    return next(
        o for o in report["objects"] if o.get("formatting") and "Hello" in o["text"]
    )


def edit(data, fmt, action, value):
    obj = target(data, fmt)
    plan = plan_repairs(
        data,
        "Sample." + fmt,
        digest(data),
        [{"object_id": obj["id"], "action": action, "value": value}],
        "formatting-test",
    )
    return apply_repairs(data, plan)


@pytest.mark.parametrize("fmt", ["docx", "pptx"])
def test_measured_contrast_and_minimal_color_repair(fmt):
    data = fixture(fmt)
    before = target(data, fmt)
    assert before["formatting"]["outcome"] == "fail"
    assert before["formatting"]["suggested_colors"] == ["000000"]
    with pytest.raises(PackageError, match="below the contrast minimum"):
        edit(data, fmt, "set_text_color", "#EEEEEE")
    output, report, receipt = edit(data, fmt, "set_text_color", "#000000")
    after = target(output, fmt)
    assert after["formatting"]["outcome"] == "pass" and after["text"] == before["text"]
    original, repaired = Package(data), Package(output)
    for name, member in original.members.items():
        if name != before["target"]["part"]:
            assert repaired.members[name] == member
    assert receipt["verification"]["operation_idempotence"] == "pass"
    node = repaired.locate(after["target"])
    prefix = "w" if fmt == "docx" else "a"
    assert len(list(node.iter(q(prefix, "r")))) == 2
    if fmt == "docx":
        assert node.find(".//w:i", NS) is not None
    else:
        assert node.find(".//a:rPr", NS).get("i") == "1"


@pytest.mark.parametrize("fmt", ["docx", "pptx"])
def test_font_size_changes_threshold_without_rewriting_text(fmt):
    data = fixture(fmt, color="888888")  # ~3.54:1: fails small text, passes large
    assert target(data, fmt)["formatting"]["outcome"] == "fail"
    output, _, _ = edit(data, fmt, "set_font_size", "18")
    assert target(output, fmt)["formatting"]["outcome"] == "pass"
    assert target(output, fmt)["formatting"]["font_sizes"] == [18]
    assert target(output, fmt)["text"] == target(data, fmt)["text"]
    assert (
        target(fixture(fmt, size=14, color="888888", bold=True), fmt)["formatting"][
            "outcome"
        ]
        == "pass"
    )
    for value in ("nan", "inf", "0", "200", "12.25", True):
        with pytest.raises(PackageError):
            edit(data, fmt, "set_font_size", value)


def test_docx_shading_highlight_and_theme_color_repair():
    data = fixture("docx", color="FFFFFF")
    pkg = Package(data)
    p = pkg.roots[pkg.main].find(".//w:p", NS)
    ppr = ET.Element(q("w", "pPr"))
    p.insert(0, ppr)
    ET.SubElement(
        ppr, q("w", "shd"), {q("w", "fill"): "000000", q("w", "val"): "clear"}
    )
    pkg.touch(pkg.main)
    data = pkg.write()
    assert target(data, "docx")["formatting"]["outcome"] == "pass"
    run = p.find("w:r", NS)
    rp = run.find("w:rPr", NS)
    ET.SubElement(rp, q("w", "highlight"), {q("w", "val"): "yellow"})
    pkg.touch(pkg.main)
    assert target(pkg.write(), "docx")["formatting"]["outcome"] == "needs_review"
    # Explicit replacement must remove theme attributes that would override RGB.
    data = fixture("docx")
    pkg = Package(data)
    p = pkg.roots[pkg.main].find(".//w:p", NS)
    rp = p.find("w:r/w:rPr", NS)
    ET.SubElement(
        rp,
        q("w", "color"),
        {q("w", "themeColor"): "accent1", q("w", "themeTint"): "80"},
    )
    pkg.touch(pkg.main)
    output, _, _ = edit(pkg.write(), "docx", "set_text_color", "#000000")
    for node in Package(output).roots[pkg.main].iter(q("w", "color")):
        assert node.get(q("w", "themeColor")) is None


def test_pptx_unknown_background_and_overlap_never_pass():
    data = fixture("pptx")
    doc = Presentation(io.BytesIO(data))
    slide = doc.slides[0]
    slide.background.fill.gradient()
    stream = io.BytesIO()
    doc.save(stream)
    assert target(stream.getvalue(), "pptx")["formatting"]["outcome"] == "needs_review"
    assert not target(stream.getvalue(), "pptx")["formatting"]["suggested_colors"]
    doc = Presentation(io.BytesIO(data))
    slide = doc.slides[0]
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = (
        "Overlap"
    )
    stream = io.BytesIO()
    doc.save(stream)
    assert target(stream.getvalue(), "pptx")["formatting"]["outcome"] == "needs_review"


@pytest.mark.parametrize("fmt", ["docx", "pptx"])
def test_ai_formatting_proposals_are_bounded_and_reviewable(monkeypatch, fmt):
    monkeypatch.setenv("OPENAI_API_KEY", "test-only")
    data = fixture(fmt)
    obj = target(data, fmt)
    proposals = [
        {
            "object_id": obj["id"],
            "action": action,
            "value": value,
            "reason": "Improve readability.",
        }
        for action, value in [
            ("set_text_color", "#EEEEEE"),
            ("set_text_color", "#000000"),
            ("set_font_size", "18"),
        ]
    ]
    with patch(
        "workshop.office.ai_drafts.chat_completion",
        return_value=json.dumps({"suggestions": proposals, "review_notes": []}),
    ) as provider:
        result = draft_repairs(data, "Sample." + fmt, digest(data))
    assert result["excluded_proposals"] == 1
    assert len(result["suggestions"]) == 2
    assert "suggested_colors" in str(provider.call_args.kwargs["messages"])
    plan = plan_repairs(
        data, "Sample." + fmt, digest(data), result["suggestions"], "ai-formatting-test"
    )
    output, _, _ = apply_repairs(data, plan)
    assert target(output, fmt)["formatting"]["outcome"] == "pass"


def test_ratio_is_not_rounded_before_comparison():
    assert ratio("000000", "FFFFFF") == 21
    assert ratio("777777", "FFFFFF") < 4.5


def test_pptx_theme_and_paragraph_inheritance_and_alpha():
    data = fixture("pptx")
    pkg = Package(data)
    part = "ppt/slides/slide1.xml"
    root = pkg.xml(part)
    for props in root.iter(q("a", "rPr")):
        props.attrib.pop("sz", None)
        for node in list(props):
            if node.tag == q("a", "solidFill"):
                props.remove(node)
    para = root.find(".//a:p", NS)
    props = ET.Element(q("a", "pPr"))
    para.insert(0, props)
    defaults = ET.SubElement(props, q("a", "defRPr"), sz="1800", b="0")
    solid = ET.SubElement(defaults, q("a", "solidFill"))
    scheme = ET.SubElement(solid, q("a", "schemeClr"), val="accent1")
    theme = pkg.xml("ppt/theme/theme1.xml")
    accent = theme.find(".//a:clrScheme/a:accent1/a:srgbClr", NS)
    accent.set("val", "888888")
    pkg.touch(part)
    pkg.touch("ppt/theme/theme1.xml")
    obj = target(pkg.write(), "pptx")
    assert obj["formatting"]["outcome"] == "pass"
    assert obj["formatting"]["font_sizes"] == [18]
    assert obj["formatting"]["foregrounds"] == ["888888"]
    ET.SubElement(scheme, q("a", "alpha"), val="50000")
    pkg.touch(part)
    assert target(pkg.write(), "pptx")["formatting"]["outcome"] == "needs_review"


def test_unknown_size_never_claims_a_marginal_ratio_failure():
    from workshop.office.formatting import sample

    assert sample("text", "888888", "FFFFFF", None, None)["outcome"] == "needs_review"
    assert sample("text", "888888", "FFFFFF", 14, None)["outcome"] == "needs_review"
    assert sample("text", "AAAAAA", "FFFFFF", None, None)["outcome"] == "fail"


def test_font_repair_of_small_text_is_advisory_not_universal_minimum():
    obj = target(fixture("docx", size=10, color="000000"), "docx")
    assert obj["formatting"]["outcome"] == "pass"
    report = OfficeEngine().inspect(
        fixture("docx", size=10, color="000000"), "Sample.docx"
    )
    row = next(f for f in report["findings"] if f["rule_id"] == "docx.text.readability")
    assert row["outcome"] == "needs_review" and "not a WCAG minimum" in row["message"]


def test_unsupported_word_text_effects_remain_unresolved():
    data = fixture("docx", color="000000")
    pkg = Package(data)
    p = pkg.roots[pkg.main].find(".//w:p", NS)
    rp = p.find("w:r/w:rPr", NS)
    ET.SubElement(rp, "{http://schemas.microsoft.com/office/word/2010/wordml}textFill")
    pkg.touch(pkg.main)
    assert target(pkg.write(), "docx")["formatting"]["outcome"] == "needs_review"

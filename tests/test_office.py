"""Portable authored corpus: object identity, conservative edits and package fidelity."""

import io
import json
import zipfile
from pathlib import Path
import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from lxml import etree as ET
from workshop.office.engine import OfficeEngine
from workshop.office.package import Package, PackageError, Conflict, NS, q
from workshop.office.repairs import plan_repairs, apply_repairs
from workshop.contracts import digest
from workshop.services import export_bundle
from web.app import app


def save(document):
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def patch(data, part, transform):
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as dest:
        for info in source.infolist():
            payload = source.read(info)
            dest.writestr(
                info, transform(payload) if info.filename == part else payload
            )
    return output.getvalue()


def docx_bytes():
    doc = Document()
    doc.add_heading("Application instructions", 0)
    doc.add_paragraph(
        "Applicant {{ users[0].name }} keeps the original template token."
    )
    for _ in range(2):
        table = doc.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, [["Name", "County"], ["Ada", "Suffolk"]]):
            for cell, value in zip(row.cells, values):
                cell.text = value
    doc.sections[0].header.paragraphs[0].text = "Header kept intact"
    return save(doc)


def pptx_bytes():
    pres = Presentation()
    for title in ["First slide", "Second slide"]:
        slide = pres.slides.add_slide(pres.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = "Readable content"
        slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(1))
    return save(pres)


def repair(data, kind, action, value, filename="Application.docx"):
    report = OfficeEngine().inspect(data, filename)
    obj = next(
        o for o in report["objects"] if o["kind"] == kind and action in o["operations"]
    )
    plan = plan_repairs(
        data,
        filename,
        digest(data),
        [{"object_id": obj["id"], "action": action, "value": value}],
        "test-retry-key",
    )
    return apply_repairs(data, plan), plan


def test_distinct_tables_and_preservation():
    data = docx_bytes()
    report = OfficeEngine().inspect(data, "Application.docx")
    rows = [f for f in report["findings"] if f["source_code"] == "WA552"]
    assert len(rows) == 2 and rows[0]["target"] != rows[1]["target"]
    (after, new, receipt), plan = repair(data, "table", "set_table_header", True)
    assert len([f for f in new["findings"] if f["source_code"] == "WA552"]) == 1
    assert "{{ users[0].name }}" in Document(io.BytesIO(after)).paragraphs[1].text
    assert receipt["changed_parts"] == ["word/document.xml"]
    assert apply_repairs(data, plan)[0] == after
    with pytest.raises(Conflict):
        apply_repairs(after, plan)
    assert new["version"]["parent_version"] == report["version"]["id"]
    assert (
        OfficeEngine().inspect(after, "Application.docx")["version"]["id"]
        == new["version"]["id"]
    )


@pytest.mark.parametrize("value", ["0", "false", "off"])
def test_false_table_header(value):
    data = docx_bytes()

    def edit(payload):
        root = ET.fromstring(payload)
        for row in root.findall(".//w:tbl/w:tr[1]", NS):
            props = ET.SubElement(row, q("w", "trPr"))
            ET.SubElement(props, q("w", "tblHeader")).set(q("w", "val"), value)
        return ET.tostring(root)

    report = OfficeEngine().inspect(
        patch(data, "word/document.xml", edit), "Application.docx"
    )
    assert len([f for f in report["findings"] if f["source_code"] == "WA552"]) == 2


def test_malformed_required_optional():
    data = docx_bytes()
    with pytest.raises(PackageError):
        OfficeEngine().inspect(
            patch(data, "word/document.xml", lambda _: b"<broken"), "Application.docx"
        )
    report = OfficeEngine().inspect(
        patch(data, "word/header1.xml", lambda _: b"<broken"), "Application.docx"
    )
    assert any(
        c.get("part") == "word/header1.xml" and c["outcome"] == "error"
        for c in report["coverage"]
    )


def test_language_and_logical_name():
    data = docx_bytes()

    def edit(payload):
        root = ET.fromstring(payload)
        for lang in root.findall(".//w:lang", NS):
            lang.getparent().remove(lang)
        style = ET.SubElement(
            root,
            q("w", "style"),
            {q("w", "styleId"): "Unused", q("w", "type"): "paragraph"},
        )
        ET.SubElement(ET.SubElement(style, q("w", "rPr")), q("w", "lang")).set(
            q("w", "val"), "en-US"
        )
        return ET.tostring(root)

    data = patch(data, "word/styles.xml", edit)
    report = OfficeEngine().inspect(data, "document.docx")
    assert any(f["source_code"] == "EA545" for f in report["findings"])
    assert any(f["source_code"] == "IA562" for f in report["findings"])
    assert not any(
        f["source_code"] == "IA562"
        for f in OfficeEngine().inspect(data, "Application for assistance.docx")[
            "findings"
        ]
    )
    (after, report, _), _ = repair(data, "language", "set_default_language", "es")
    assert not any(f["source_code"] == "EA545" for f in report["findings"])
    Document(io.BytesIO(after))


def test_title_export_and_hash():
    data = docx_bytes()
    (after, report, receipt), _ = repair(
        data, "metadata", "set_title", "Housing application"
    )
    assert Document(io.BytesIO(after)).core_properties.title == "Housing application"
    bundle = export_bundle(
        after,
        "Application.docx",
        {
            "source_sha256": digest(after),
            "version": report["version"],
            "repair_receipts": [receipt],
        },
    )
    with zipfile.ZipFile(io.BytesIO(bundle)) as archive:
        manifest = json.loads(archive.read("review-manifest.json"))
        assert archive.read("Application.docx") == after
        assert manifest["document"] == report["version"]
    with pytest.raises(PackageError):
        export_bundle(after, "Application.docx", {"source_sha256": digest(data)})


def test_presentation_native_order_title_and_preservation():
    data = pptx_bytes()

    def reverse(payload):
        root = ET.fromstring(payload)
        lst = root.find(q("p", "sldIdLst"))
        lst[:] = list(reversed(lst))
        return ET.tostring(root)

    data = patch(data, "ppt/presentation.xml", reverse)
    report = OfficeEngine().inspect(data, "Training.pptx")
    titles = [o for o in report["objects"] if "set_existing_title" in o["operations"]]
    assert [o["text"] for o in titles] == ["Second slide", "First slide"]
    decision = {
        "object_id": titles[0]["id"],
        "action": "set_existing_title",
        "value": "Accessible training",
    }
    plan = plan_repairs(
        data, "Training.pptx", digest(data), [decision], "title-change-1"
    )
    after, new, receipt = apply_repairs(data, plan)
    assert (
        Presentation(io.BytesIO(after)).slides[0].shapes.title.text
        == "Accessible training"
    )
    assert receipt["changed_parts"] == ["ppt/slides/slide2.xml"]
    assert any(
        f["rule_id"] == "pptx.text.contrast" and f["outcome"] == "needs_review"
        for f in new["findings"]
    )


def test_api_stale_and_bounded_errors():
    client = app.test_client()
    data = docx_bytes()
    result = client.post(
        "/api/documents/inspect", data={"file": (io.BytesIO(data), "Housing.docx")}
    )
    assert result.status_code == 200
    report = result.json["data"]
    obj = next(o for o in report["objects"] if o["kind"] == "metadata")
    result = client.post(
        "/api/documents/plan",
        data={
            "file": (io.BytesIO(data), "Housing.docx"),
            "expected_sha256": "old",
            "decisions": json.dumps(
                [{"object_id": obj["id"], "action": "set_title", "value": "Housing"}]
            ),
            "idempotency_key": "test-plan-key",
        },
    )
    assert result.status_code == 409
    assert (
        client.post(
            "/api/documents/inspect", data={"file": (io.BytesIO(b"bad"), "bad.docx")}
        ).status_code
        == 400
    )


def test_visual_description_decorative_and_multilingual_preservation(tmp_path):
    from PIL import Image

    image = tmp_path / "picture.png"
    Image.new("RGB", (12, 12), "blue").save(image)
    document = Document()
    paragraph = document.add_paragraph("English text ")
    run = paragraph.add_run("Bonjour")
    language = ET.SubElement(run._r.get_or_add_rPr(), q("w", "lang"))
    language.set(q("w", "val"), "fr-FR")
    document.add_picture(str(image))
    data = save(document)
    (after, report, _), _ = repair(
        data, "image", "set_alt_text", "Blue square used as a legend key"
    )
    assert not any(f["source_code"] == "EA541" for f in report["findings"])
    (decorated, report, _), _ = repair(after, "image", "mark_decorative", "")
    assert next(o for o in report["objects"] if o["kind"] == "image")["decorative"]
    assert not any(f["source_code"] == "WA547" for f in report["findings"])
    (after, _, _), _ = repair(decorated, "language", "set_default_language", "en-US")
    root = Package(after).roots["word/document.xml"]
    assert root.xpath("//w:rPr/w:lang/@w:val", namespaces=NS) == ["fr-FR"]


def test_pptx_hidden_title_language_and_image(tmp_path):
    from PIL import Image

    image = tmp_path / "picture.png"
    Image.new("RGB", (12, 12), "blue").save(image)
    pres = Presentation()
    slide = pres.slides.add_slide(pres.slide_layouts[1])
    slide.shapes.title.text = "Hidden but meaningful"
    slide.shapes.title.left = -Inches(8)
    slide.placeholders[1].text = "Body text"
    slide.shapes.add_picture(str(image), Inches(1), Inches(2))
    data = save(pres)
    report = OfficeEngine().inspect(data, "Workshop.pptx")
    assert any(
        f["rule_id"] == "pptx.slide.title" and f["outcome"] == "pass"
        for f in report["findings"]
    )
    (after, report, _), _ = repair(
        data, "image", "set_alt_text", "Blue square legend key", "Workshop.pptx"
    )
    assert not any(
        f["rule_id"] == "pptx.visual.alternative" and f["outcome"] == "fail"
        for f in report["findings"]
    )
    (after, report, _), _ = repair(
        after, "language", "set_default_language", "es", "Workshop.pptx"
    )
    assert not any(
        f["rule_id"] == "pptx.text.language" and f["outcome"] != "pass"
        for f in report["findings"]
    )
    assert Presentation(io.BytesIO(after)).slides[0].shapes.title.left < 0


def test_ordinary_link_edit_preserves_destination():
    document = Document()
    paragraph = document.add_paragraph()
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    relation = paragraph.part.relate_to(
        "https://example.org/application", RT.HYPERLINK, is_external=True
    )
    hyperlink = ET.SubElement(
        paragraph._p, q("w", "hyperlink"), {q("r", "id"): relation}
    )
    run = ET.SubElement(hyperlink, q("w", "r"))
    ET.SubElement(run, q("w", "t")).text = "Click here"
    data = save(document)
    (after, _, receipt), _ = repair(
        data, "link", "set_link_text", "Housing application"
    )
    package = Package(after)
    assert "Housing application" in "".join(
        package.roots["word/document.xml"].itertext()
    )
    assert (
        package.members["word/_rels/document.xml.rels"]
        == Package(data).members["word/_rels/document.xml.rels"]
    )
    assert receipt["verification"]["operation_idempotence"] == "pass"


def test_signed_pdf_cannot_be_rewritten():
    import pikepdf

    source = Path("tests/plain-document.pdf").read_bytes()
    with pikepdf.open(io.BytesIO(source)) as pdf:
        signature = pdf.make_indirect(
            pikepdf.Dictionary(Type=pikepdf.Name("/Sig"), ByteRange=[0, 1, 2, 3])
        )
        pdf.Root.Perms = pikepdf.Dictionary(DocMDP=signature)
        stream = io.BytesIO()
        pdf.save(stream)
    source = stream.getvalue()
    client = app.test_client()
    inspection = client.post(
        "/pdf-labeler/api/accessibility-inspect",
        data={"file": (io.BytesIO(source), "signed.pdf")},
    )
    assert inspection.json["data"]["edit_policy"]["editable"] is False
    repair = client.post(
        "/pdf-labeler/api/accessibility-remediate",
        data={
            "file": (io.BytesIO(source), "signed.pdf"),
            "action": "metadata",
            "title": "changed",
        },
    )
    assert (
        repair.status_code == 400
        and "inspection-only" in repair.json["error"]["message"]
    )
